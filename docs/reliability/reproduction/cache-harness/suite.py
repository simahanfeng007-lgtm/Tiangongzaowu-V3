"""Synthetic benchmark tasks and artifact-only checks outside the product gate."""
from pathlib import Path
import csv
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET

BASE = Path(__file__).resolve().parent
COMMON = "所有操作限定在当前工作区，禁止访问桌面。请实际生成并检查交付文件。"

TASKS = [
 {"id":"files","family":"文件整理","prompt":"将 inbox 下所有文件复制到 organized，保留各自相对于 inbox 的子目录和原始字节。生成 manifest.csv，表头 source,target,size,sha256，每个文件一行；source 以 inbox/ 开头，target 以 organized/ 开头。保留输入文件。"},
 {"id":"data","family":"数据汇总","prompt":"根据 orders.csv 汇总。只统计 status=paid 的行，退款行 refunded 不计入。生成 summary.json，顶层字段 regions、grand_total_cents；regions 按地区名称升序排列，每项字段 region、quantity、revenue_cents，收入为 quantity 乘 unit_price_cents。全部金额用整数分。保留输入。"},
 {"id":"code","family":"代码修复","prompt":"修复 calc.py 中 net_total(rows) 的计算错误：忽略 cancelled=true 的记录，其他记录按 qty*unit_cents-discount_cents 求和，结果为整数分。保留命令行接口 python calc.py input.json output.json，输出 {\"total_cents\":整数}。实际运行样例 orders.json，生成 result.json。不要修改 orders.json。"},
 {"id":"docx","family":"Word文档","prompt":"读取 notice.json，生成 notice.docx，包含源数据给定的标题、日期、时间、地点、主持人、两项议程及注意事项；姓名和时间逐字保留，不添加虚构事项。文档短而清楚即可，保留输入文件。"},
 {"id":"xlsx","family":"Excel表格","prompt":"根据 products.csv 生成 sales.xlsx，工作表名为 Sales。A1:D1 依次为 product,quantity,unit_price_cents,revenue_cents；按输入顺序写三行，D2:D4 使用同一行 B 乘 C 的 Excel 公式。A5 写 TOTAL，D5 使用 SUM(D2:D4) 公式。保留输入，不要求公式缓存值。"},
 {"id":"pptx","family":"PPT演示","prompt":"读取 brief.json，生成 briefing.pptx，恰好四页。每页标题按源文件 slides 顺序逐字使用，正文包含对应页给定的全部 facts，不改写数字、单位和姓名。无需商务风格评分或额外 CTA。保留输入。"},
 {"id":"story","family":"短篇写作","prompt":"读取 story_brief.json，按其中人物和关键事实写 story.md。三个章节标题依次为“一、灯塔”“二、来信”“三、归航”。全文含标题共 400 至 700 个汉字，必须出现两个人物的完整姓名、船只编号和约定日期；情节完整、有结尾。保留输入文件。"},
 {"id":"image","family":"图片处理","prompt":"将 source.png 中心裁剪为正方形，再缩放成 256×256 像素，保存 thumbnail.png。保留颜色和输入文件，不加文字、边框或其他图案。"},
]

def json_bytes(value):
    return (json.dumps(value,ensure_ascii=False,indent=2)+"\n").encode()

def prepare():
    from PIL import Image
    fixture_root=BASE/'fixtures'
    fixture_root.mkdir(exist_ok=True)
    fixtures={
      'files':{'inbox/a/notes.txt':'甲组记录\n'.encode(),'inbox/b/notes.txt':'乙组记录\n'.encode(),
               'inbox/data.csv':b'id,value\n1,17\n2,25\n','inbox/raw.bin':bytes(range(256)),'inbox/empty.txt':b''},
      'data':{'orders.csv':b'region,quantity,unit_price_cents,status\nEast,3,1299,paid\nWest,2,850,paid\nEast,4,99,paid\nNorth,1,5000,refunded\nWest,0,777,paid\nNorth,2,1250,paid\n'},
      'code':{'calc.py':b'import json,sys\ndef net_total(rows):\n    return sum(r["qty"]*r["unit_cents"] for r in rows)\nif __name__ == "__main__":\n    rows=json.load(open(sys.argv[1]))\n    json.dump({"total_cents":net_total(rows)},open(sys.argv[2],"w"))\n',
              'orders.json':json_bytes([{'qty':2,'unit_cents':1250,'discount_cents':100,'cancelled':False},{'qty':3,'unit_cents':99,'discount_cents':0,'cancelled':True},{'qty':1,'unit_cents':650,'discount_cents':50,'cancelled':False}])},
      'docx':{'notice.json':json_bytes({'title':'联合项目例会通知','date':'2026年10月8日','time':'09:30','location':'青松会议室','host':'林舟','agenda':['核对第三季度订单','确定下月交付时间'],'note':'请携带原始记录'})},
      'xlsx':{'products.csv':'product,quantity,unit_price_cents\n铅笔,3,120\n笔记本,2,850\n文件夹,4,299\n'.encode()},
      'pptx':{'brief.json':json_bytes({'slides':[{'title':'项目概况','facts':['海岬计划','负责人：许遥']},{'title':'本周进展','facts':['完成17项','剩余8项']},{'title':'资源安排','facts':['预算4200元','参与人数6人']},{'title':'下一步','facts':['10月8日复核','保留2天缓冲']}]})},
      'story':{'story_brief.json':json_bytes({'characters':['林舟','许遥'],'boat':'R-17','date':'10月8日','premise':'停用多年的灯塔再次亮起，两人在旧信中发现约定，最终驾船归航。'})},
      'image':{},
      'smoke':{'input.json':json_bytes({'left':17,'right':25,'label':'dictionary codec smoke'})},
    }
    im=Image.new('RGB',(320,160),(240,30,30))
    for x in range(160,320):
        for y in range(160):im.putpixel((x,y),(20,70,230))
    b=io.BytesIO();im.save(b,format='PNG');fixtures['image']['source.png']=b.getvalue()
    for task,files in fixtures.items():
        for name,data in files.items():
            p=fixture_root/task/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
    tasks=[{**t,'prompt':t['prompt']+COMMON} for t in TASKS]
    manifest={'schema':'dictionary.reconstructed-paired-experiment.v1','source_commit':'a372dbc7fb7d5218c07d96c33d21d29da4f48694',
      'arms':{'A':'Unmodified normal main runtime + symmetric telemetry','C':'Same runtime + complete lossless primary dictionary context + action short-code decoder'},
      'replicates':2,'task_deadline_seconds':240,'pilot_deadline_seconds':180,
      'tasks':tasks,'fixtures_sha256':{t:{n:hashlib.sha256(v).hexdigest() for n,v in fs.items()} for t,fs in fixtures.items()},
      'pairing':'One task and replicate per pair, A/C concurrent isolated processes. Launch order alternates. No state or learned memories shared.',
      'scope':'Reconstructed experiment, not reproduction of unavailable old implementation; one configured model, eight synthetic local tasks.',
      'primary_endpoint':'Independent task check passes AND request status COMPLETED before deadline',
      'secondary_endpoints':['artifact-only check','latency','prompt/completion/cache tokens','model wire requests','schema discovery','registered compositions','short-code use','step failure counts','false completion','output correct but runtime incomplete'],
      'artifact_checks':'Only requirements explicitly in each frozen prompt; no invented business/style gates; story check covers mechanical constraints, not literary quality.',
      'attribution_limit':'A vs C tests the combined full-context and short-code package; it cannot identify each component causal contribution.'}
    (BASE/'experiment.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    return manifest

def check(task, workspace):
    w=Path(workspace);failures=[];details={}
    def require(ok,label):
        if not ok:failures.append(label)
    definition=json.loads((BASE/'experiment.json').read_text())
    for name,h in definition['fixtures_sha256'][task].items():
        if task=='code' and name=='calc.py':continue
        p=w/name;require(p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()==h,'input_changed:'+name)
    try:
        if task=='smoke':
            require(json.loads((w/'result.json').read_text())=={'sum':42,'label':'dictionary codec smoke'},'json_content')
        elif task=='files':
            with (w/'manifest.csv').open(encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f);rows=list(reader)
                require(reader.fieldnames==['source','target','size','sha256'],'manifest_headers')
            expected=definition['fixtures_sha256']['files'];require(len(rows)==len(expected),'manifest_rows')
            by_source={r['source'].replace('\\','/'):r for r in rows};require(set(by_source)==set(expected),'manifest_sources')
            for source,h in expected.items():
                target='organized/'+source.removeprefix('inbox/');p=w/target
                require(p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()==h,'copy_bytes:'+source)
                r=by_source.get(source,{})
                require(r.get('target','').replace('\\','/')==target and r.get('sha256')==h and int(r.get('size',-1))==(w/source).stat().st_size,'manifest_values:'+source)
        elif task=='data':
            expected={'regions':[{'region':'East','quantity':7,'revenue_cents':4293},{'region':'North','quantity':2,'revenue_cents':2500},{'region':'West','quantity':2,'revenue_cents':1700}],'grand_total_cents':8493}
            require(json.loads((w/'summary.json').read_text())==expected,'aggregates')
        elif task=='code':
            require(json.loads((w/'result.json').read_text())=={'total_cents':3000},'sample_result')
            cases=[[],[{'qty':2,'unit_cents':101,'discount_cents':3,'cancelled':False}],
                   [{'qty':9,'unit_cents':99,'discount_cents':8,'cancelled':True}],
                   [{'qty':-1,'unit_cents':123,'discount_cents':7,'cancelled':False},{'qty':3,'unit_cents':10,'discount_cents':1,'cancelled':False}]]
            script='import importlib.util,json,sys; s=importlib.util.spec_from_file_location("tested",sys.argv[1]); m=importlib.util.module_from_spec(s); s.loader.exec_module(m); print(json.dumps([m.net_total(x) for x in json.loads(sys.stdin.read())]))'
            run=subprocess.run(['bwrap','--unshare-net','--die-with-parent','--ro-bind','/','/','--tmpfs','/tmp','--ro-bind',str(w),str(w),'--chdir',str(w),sys.executable,'-I','-B','-c',script,str(w/'calc.py')],input=json.dumps(cases),text=True,capture_output=True,timeout=15,env={'PATH':os.environ['PATH'],'HOME':'/tmp'})
            require(run.returncode==0 and json.loads(run.stdout)==[0,199,0,-101],'independent_code_cases')
            details['hidden_cases']=4
        elif task=='docx':
            with zipfile.ZipFile(w/'notice.docx') as z:text=''.join(ET.fromstring(z.read('word/document.xml')).itertext())
            source=json.loads((w/'notice.json').read_text())
            for value in [source[k] for k in ('title','date','time','location','host','note')]+source['agenda']:require(value in text,'missing_text:'+value)
        elif task=='xlsx':
            import openpyxl
            wb=openpyxl.load_workbook(w/'sales.xlsx',data_only=False);sh=wb['Sales']
            require([sh.cell(1,c).value for c in range(1,5)]==['product','quantity','unit_price_cents','revenue_cents'],'headers')
            expected=[['铅笔',3,120],['笔记本',2,850],['文件夹',4,299]]
            for i,row in enumerate(expected,2):
                require([sh.cell(i,c).value for c in range(1,4)]==row,'data_row:'+str(i))
                require(str(sh.cell(i,4).value).replace(' ','').upper() in (f'=B{i}*C{i}',f'=C{i}*B{i}'),'row_formula:'+str(i))
            require(sh['A5'].value=='TOTAL' and str(sh['D5'].value).replace(' ','').upper()=='=SUM(D2:D4)','total_formula')
        elif task=='pptx':
            with zipfile.ZipFile(w/'briefing.pptx') as z:
                names=sorted((x for x in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',x)),key=lambda x:int(re.search(r'(\d+)\.xml',x)[1]))
                texts=[''.join(ET.fromstring(z.read(n)).itertext()) for n in names]
            expected=json.loads((w/'brief.json').read_text())['slides'];require(len(texts)==4,'slide_count')
            for i,row in enumerate(expected):
                for value in [row['title']]+row['facts']:require(i<len(texts) and value in texts[i],'slide_text:'+value)
        elif task=='story':
            text=(w/'story.md').read_text();source=json.loads((w/'story_brief.json').read_text())
            titles=['一、灯塔','二、来信','三、归航'];positions=[text.find(x) for x in titles]
            require(all(x>=0 for x in positions) and positions==sorted(positions),'chapter_order')
            for value in source['characters']+[source['boat']]:require(value in text,'story_fact:'+value)
            date_forms={'10月8日':('10月8日','十月八日','10月八日','十月8日')}
            require(any(value in text for value in date_forms.get(source['date'],(source['date'],))), 'story_fact:'+source['date'])
            count=len(re.findall(r'[\u4e00-\u9fff]',text));details['cjk_characters']=count
            require(400<=count<=700,'explicit_character_range')
        elif task=='image':
            from PIL import Image
            im=Image.open(w/'thumbnail.png').convert('RGB');require(im.size==(256,256),'image_size')
            from PIL import ImageChops, ImageStat
            source=Image.open(BASE/'fixtures/image/source.png').convert('RGB')
            expected=source.crop((80,0,240,160)).resize((256,256),Image.Resampling.LANCZOS)
            if im.size==(256,256):
                delta=ImageChops.difference(im,expected);mae=sum(ImageStat.Stat(delta).mean)/3
                maximum=max(high for low,high in delta.getextrema())
                details.update(pixel_mae=mae,max_channel_error=maximum)
                require(mae<=2 and maximum<=8,'center_crop_pixels')
            details['size']=list(im.size)
        else:raise ValueError('unknown task')
    except Exception as exc:
        failures.append(type(exc).__name__+':'+str(exc)[:250])
    return {'pass':not failures,'failures':failures,'details':details}

if __name__=='__main__':
    print(json.dumps({'tasks':len(prepare()['tasks']),'status':'frozen'},ensure_ascii=False))
