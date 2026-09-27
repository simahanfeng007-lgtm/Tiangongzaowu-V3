"""Fresh ontology-only real model acceptance. All attempts kept separately."""
from pathlib import Path
import argparse
import collections
import hashlib
import json
import os
import secrets
import signal
import socket
import sqlite3
import subprocess
import sys
import time
import urllib.request
import urllib.error

BASE = Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('case', choices=['local', 'browser', 'mcp', 'quality', 'document', 'novel', 'local-apps', 'media', 'mcp-tasks'])
parser.add_argument('--attempt', default='1')
parser.add_argument('--source-root', type=Path, default=BASE.parents[1])
parser.add_argument('--output-root', type=Path, required=True)
parser.add_argument('--model-config', type=Path, required=True)
parser.add_argument('--vision-provider')
args = parser.parse_args()
if os.name != 'posix':
    parser.error('This development harness uses POSIX process groups; native Windows acceptance is separate.')
if not args.attempt or any(c not in 'abcdefghijklmnopqrstuvwxyz0123456789-' for c in args.attempt):
    parser.error('attempt must contain lowercase letters, digits or hyphens')
SOURCE = args.source_root.resolve()
CONFIG = args.model_config.resolve(strict=True)
sys.path.insert(0, str(SOURCE / 'src'))
from total_gateway.release_manifest import generate_release_manifest, release_manifest_bytes

os.umask(0o077)
out = args.output_root.resolve() / 'live' / (args.case + '-' + args.attempt)
out.mkdir(parents=True, exist_ok=False)
for name in ['profile/.tiangong/v3','profile/.tiangong-v3','workspace','state','life','world']:
    (out / name).mkdir(parents=True, exist_ok=True)
(out / 'profile/.tiangong/api_keys.json').symlink_to(CONFIG)
workspace = out / 'workspace'
if args.case == 'local-apps':
    with sqlite3.connect(workspace/'records.sqlite') as db:
        db.execute('CREATE TABLE records(name TEXT,value INTEGER)')
        db.executemany('INSERT INTO records VALUES(?,?)', [('row-'+str(i),i) for i in range(235)])
    (workspace/'bad.csv').write_text('name,value\nnew,10\nbroken,11,extra\n')
    (workspace/'good.csv').write_text('name,value\nnew,10\n')
    prompt = '对工作区 records.sqlite 实际使用 sqlite.schema.read，先尝试用 sqlite.table.import_csv 导入 bad.csv 并确认失败后没有部分写入，再以 good.csv 修复该次导入，关联 repair_of。随后用 sqlite.table.export_csv 全量导出 records 表到 all.csv，使用 sqlite.backup.create 备份到 backup.sqlite，并回读核对数据库、CSV 和备份的总行数与 value 总和。用 obsidian.note.create 在 vault/结果.md 记录事实，再通过 note.read 取得版本、note.update 增加失败及恢复记录、link.create 加上引用核对记录；使用 vault.list、search、graph.export 实际检查并导出 links.json。报告 report.md 包含实际结果并回读。只修改 records.sqlite 和新生成文件，不修改两个输入 CSV；复用上述专用接口，不用通用写文件代替它们。'
elif args.case == 'media':
    from PIL import Image, ImageDraw
    picture=Image.new('RGB',(600,260),'white');draw=ImageDraw.Draw(picture)
    for x in (20,140,260): draw.rectangle((x,40,x+70,110),fill='blue')
    draw.polygon([(440,40),(390,130),(490,130)],fill='red')
    picture.save(workspace/'unknown.png')
    prompt = '请用 image.observe 实际观察工作区 unknown.png。识别蓝色正方形和红色三角形各有几个，将两个整数存为 result.json 的 blue_squares、red_triangles 字段，保存包含观察依据、源文件版本和观察范围的 report.md，并实际重新读取 JSON 和报告。不要用文件名、元数据或自己写像素检测程序代替视觉模型观察；不要修改图片。'
elif args.case == 'local':
    (workspace / 'input.csv').write_text('city,amount\n杭州,7\n上海,8\n杭州,5\n上海,-2\n杭州,0\n', encoding='utf-8-sig')
    (workspace / 'rules.json').write_text('{"output":"totals.json","report":"report.md","operation":"sum amount per city"}')
    (workspace / 'aggregate.py').write_text("import csv,json\nr={}\nfor row in csv.DictReader(open('input.csv',encoding='utf-8-sig')):r[row['city']]=int(row['amount'])\njson.dump(r,open('totals.json','w'),ensure_ascii=False)\n")
    prompt = '读取工作区 input.csv、rules.json 和现有 aggregate.py。现有脚本可能有错误，请实际运行核对，修复聚合逻辑并重新运行，生成按城市求和的 totals.json 和含输入行数、各城市合计、总计及处理方法的中文 report.md。必须重新读取产物与原始数据核对，不改变原始输入。使用本体组合和隔离执行；无需访问桌面。'
elif args.case == 'browser':
    (workspace / 'page.html').write_text('<!doctype html><html><body><h1>真实交互验收</h1><div id="status">尚未点击</div><button id="advance" onclick="document.querySelector(\'#status\').textContent=\'第二页 合计=73\';document.querySelector(\'#download\').hidden=false">下一页</button><a hidden id="download" href="data:application/json,%7B%22total%22%3A73%7D">下载结果</a></body></html>')
    prompt = '请使用真实浏览器打开工作区 page.html，点击 #advance。读取点击后的页面文字，记录第二页合计；找到 #download 的真实链接并下载为 result.json，回读下载内容，保存真实页面截图和简短中文 report.md。不要改写源网页，不把静态抓取当成浏览器交互。'
elif args.case == 'document':
    from docx import Document
    document = Document()
    document.add_paragraph('库存盘点记录')
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [('商品','数量'), ('SKU-A','7'), ('SKU-B','11')]):
        for cell, value in zip(row.cells, values): cell.text = value
    document.sections[0].header.paragraphs[0].text = '页眉不属于正文提取范围'
    document.save(workspace/'inventory.docx')
    (workspace/'corrupt.docx').write_bytes(b'not a DOCX archive')
    prompt = '用 qc.docx.delivery_check 分别实际观察 inventory.docx 与 corrupt.docx。根据可读取正文和表格，写 report.md，列出 SKU-A、SKU-B 数量及合计，并如实记录损坏文件的读取失败。只要求正文和表格内容，不要求版式或页眉；禁止修改原始文件，报告写完后重新读取确认。'
elif args.case == 'novel':
    sys.path.insert(0, str(SOURCE / 'app/backend/tiangong-backend'))
    from v3.novel_system import NovelSystemEngine
    engine = NovelSystemEngine(workspace/'managed-novel')
    engine.create_project({'title':'星光记录','genre':'微型科幻','planned_chapters':1,'target_words':10000})
    for section, data in (
        ('story', {'soul':'保留指定原文','protected_anchors':[]}),
        ('world', {'rules':['由作者定义故事世界']}), ('characters', [{'id':'c1','name':'旅人','initial':{'alive':True,'location':'l1'}}]),
        ('calendar', {'ticks_per_year':365,'start_tick':0}), ('locations', [{'id':'l1','name':'站台'}]),
        ('plot_events', [{'id':'e1','chapter':1,'participants':['c1'],'location':'l1','start_tick':0,'duration_ticks':1}]),
        ('chapters', [{'number':1,'title':'星光','event_ids':['e1']}]), ('settings', {'min_chapter_chars':2500})):
        engine.update_blueprint({'section':section,'data':data})
    engine.compile_blueprint({})
    lease=engine.checkout_chapter({'chapter_number':1})
    engine.submit_chapter({'lease_id':lease['lease_id'],'chapter_number':1,'title':'星光','content':'星光亮起。','actual':{}})
    (workspace/'novel-request.json').write_text(json.dumps({'title':'星光','content':'星光熄灭。','actual':{}},ensure_ascii=False))
    prompt = '工作区 managed-novel 是已经记录一章的测试小说项目，现有正文“星光亮起。”需要实际修正。本任务只要求原样改为 novel-request.json 指定的短正文“星光熄灭。”，标题星光，actual 使用空对象；不要扩写，不要求兑现旧计划事件或达到配置字数。请实际读取现有正文和 novel.context.query 的 chapter 记录，使用当前记录的 sha256 作为 novel.chapter.checkout 的 revision_of，再通过 novel.chapter.submit 修订末章；actual 必须为空，不修改历史故事注释。重新读取写入的章节正文，运行项目审计，将事务状态、字节和内容观察范围写入 report.md 并回读。工具 committed/章节计数本身不代表内容获批，由最终对抗裁判判断本次原样记录任务是否完成。禁止用通用写文件绕过小说事务。'
elif args.case == 'quality':
    (workspace/'one.py').write_text('value = 18\n')
    (workspace/'test_marker.py').write_text("from pathlib import Path\nPath('must-not-execute.txt').write_text('executed')\n")
    prompt = '仅做静态观察：用 qc.code.delivery_check 检查工作区代码，读取检查结果后写 report.md，报告实际语法检查范围和是否执行了测试。禁止执行 one.py、test_marker.py、pytest 或其他项目代码；不要用代码行数、文件名或缺 README 判定业务是否完成，不需要补写代码或测试。'
else:
    if args.case == 'mcp-tasks':
        from task_fixture import start
    else:
        from mcp_fixture import start
    fixture = start(out)
    cfg={'servers':{'ontology-fixture':{'transport':'streamable_http',
         'url':f'http://127.0.0.1:{fixture.server_port}/mcp','applications':['core.filesystem'],
         'environment':{'location':'controlled local MCP test service','workspace':'test records','account_label':'isolated fixture'}}}}
    (out/'profile/.tiangong/v3/mcp_servers.json').write_text(json.dumps(cfg))
    prompt = '请通过现有 MCP 发现关联 core.filesystem 的测试环境，连接后分页读取实际工具契约。使用 ontology-fixture 创建且只创建一条 label 为 本体连接验收、value 为 41 的记录，再通过服务读取该记录的实际内容与版本，生成中文 report.md。不要猜工具名称、参数或对象 ID，不把本机受控服务说成真实云端厂商应用。'
    if args.case == 'mcp-tasks':
        prompt = '请通过 MCP 发现 ontology-fixture 的真实工具和任务能力，保留并在后续动作绑定实际 connection_id。在该受控服务异步创建一条 label 为 本体连接验收、value 为 41、delay_seconds 为 1 的记录，任务受理后记录 taskId 和 connection_fingerprint，通过 tasks.get/result 查询完成结果，再实际调用读取工具回读对象。另建一条 label 为 取消验收、value 为 0、delay_seconds 为 300 的异步任务，实际取消并重新查询确认取消状态；不要等待其执行，不重发任何创建。使用 tasks.list 核对两个任务并关闭本任务会话。写 report.md 并回读，如实区分受理、完成、取消和目标回读，不把受控服务当真实厂商。'

settings = json.loads(CONFIG.read_text())
provider = settings['_default_provider']
profile = settings['_provider_inputs'][provider]
key = settings.get('_api_keys', {}).get(provider) or profile.get('api_key')
if not key:
    raise SystemExit('The selected profile must have its credential in the supplied private model configuration.')
with socket.socket() as sock:
    sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
url = f'http://127.0.0.1:{port}'
env = {k:v for k,v in os.environ.items() if not k.startswith('TIANGONG_') and k not in ('PYTHONPATH','PYTHONSTARTUP')}
env.update({'PYTHONDONTWRITEBYTECODE':'1','ONTOLOGY_RUN':str(out),'ONTOLOGY_SOURCE':str(SOURCE),
    'ONTOLOGY_ACCEPTANCE_KEY':key,'TIANGONG_GATEWAY_ENVIRONMENT':'test',
    'TIANGONG_GATEWAY_DEPLOYMENT_MODE':'embedded','TIANGONG_GATEWAY_PORT':str(port),'TIANGONG_GATEWAY_URL':url,
    'TIANGONG_GATEWAY_STATE_ROOT':str(out/'state'),'TIANGONG_GATEWAY_WORKSPACE_ROOT':str(workspace),
    'TIANGONG_GATEWAY_RELEASE_SOURCE_ROOT':str(SOURCE),'TIANGONG_GATEWAY_RELEASE_MANIFEST_PATH':str(out/'release-manifest.json'),
    'TIANGONG_LIFE_DATA_ROOT':str(out/'life'),'TIANGONG_LIFE_RUNTIME_ROOT':str(out/'complete-life'),
    'TIANGONG_RUN_STATE_DIR':str(out/'state'),'TIANGONG_V3_STATE_DIR':str(out/'state'),
    'TIANGONG_WORLD_STATE_ROOT':str(out/'world'),
    'TIANGONG_OMNI_BODY_ROOT':str(SOURCE/'app/backend/tiangong-backend/_internal/omni_body_skill'),
    'TIANGONG_OMNI_BODY_STATE_ROOT':str(out/'omni-body'),'TIANGONG_OMNI_BODY_WORKSPACE':str(workspace)})
if args.vision_provider:
    if args.vision_provider not in settings['_provider_inputs']:
        raise ValueError('The selected vision provider must be configured in the supplied private profile.')
    env['TIANGONG_VISION_PROVIDER']=args.vision_provider
for field in ['TIANGONG_BACKEND_INTERNAL_TOKEN','TIANGONG_LIFE_INTERNAL_TOKEN','TIANGONG_GATEWAY_COMMUNICATION_TOKEN',
              'TIANGONG_GATEWAY_LIFE_INTENT_TOKEN','TIANGONG_GATEWAY_SHADOW_TOKEN','TIANGONG_DESKTOP_TOKEN']:
    env[field] = secrets.token_hex(32)
(out/'release-manifest.json').write_bytes(release_manifest_bytes(generate_release_manifest(SOURCE)))
tracked = subprocess.check_output(['git','ls-files','-z'],cwd=SOURCE).split(b'\0')
identity = {os.fsdecode(name):hashlib.sha256((SOURCE/os.fsdecode(name)).read_bytes()).hexdigest()
            for name in tracked if name and (SOURCE/os.fsdecode(name)).is_file()}
(out/'source-files.json').write_text(json.dumps(identity, sort_keys=True))
record = {'case':args.case,'prompt':prompt,'scope':'MODEL_LIVE_LOCAL','configured_provider':provider,
    'configured_model':profile['model_name'],'source_identity':'working_tree_file_hashes',
    'base_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=SOURCE,text=True).strip(),
    'source_files_sha256':hashlib.sha256((out/'source-files.json').read_bytes()).hexdigest(),
    'driver_files_sha256':{n:hashlib.sha256((BASE/n).read_bytes()).hexdigest() for n in ('run.py','gateway_entry.py','instrumentation.py','mcp_fixture.py','task_fixture.py')},
    'input_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in workspace.iterdir() if p.is_file()},
    'initial_managed_fixture_sha256':{p.relative_to(workspace).as_posix():hashlib.sha256(p.read_bytes()).hexdigest()
        for p in (workspace/'managed-novel').rglob('*') if p.is_file()},
    'started_at':time.strftime('%Y-%m-%dT%H:%M:%S%z')}
def save():
    (out/'result.json').write_text(json.dumps(record, ensure_ascii=False, indent=2))
def api(path, body=None):
    req = urllib.request.Request(url+path, data=json.dumps(body).encode() if body is not None else None,
          headers={'Content-Type':'application/json','X-Tiangong-Token':env['TIANGONG_DESKTOP_TOKEN']})
    try:
        with urllib.request.urlopen(req,timeout=15) as r: return {'status':r.status,'body':json.load(r)}
    except urllib.error.HTTPError as e:
        return {'status':e.code,'body':json.load(e)}
save()
with (out/'gateway.log').open('w') as log:
    process = subprocess.Popen([sys.executable,'-B',str(BASE/'gateway_entry.py')],cwd=out,env=env,
                               stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
    try:
        for _ in range(45):
            if process.poll() is not None: break
            try:
                ready=api('/ready')
                if ready['status']==200: break
            except (OSError,ValueError): pass
            time.sleep(1)
        record['inbound']=api('/api/v1/gateway/desktop/inbound',{
            'presentation_request_id':'ontology-'+secrets.token_hex(8),'session_id':'ontology-'+secrets.token_hex(8),
            'message_id':'message-'+secrets.token_hex(8),'submitted_at_ms':int(time.time()*1000),
            'text':prompt,'attachments':[]})
        rid=record['inbound']['body']['gateway_request_id'];record['request_id']=rid;save()
        started=time.monotonic();history=[]
        while time.monotonic()-started < 600:
            status=api('/api/v1/gateway/desktop/status?request_id='+rid)
            if not history or history[-1]!=status:
                history.append(status);(out/'status-history.json').write_text(json.dumps(history,ensure_ascii=False,indent=2))
            record['final_status']=status;record['request_status']=status.get('body',{}).get('run',{}).get('status')
            if record['request_status'] in ('COMPLETED','FAILED','CANCELLED','STOPPED'): break
            time.sleep(1)
        else:
            record['deadline_exceeded']=True
            record['cancel_response']=api('/api/v1/run/control',{'action':'cancel','request_id':rid})
        record['wall_seconds']=round(time.monotonic()-started,3)
    except Exception as e:
        record['driver_error']=type(e).__name__+':'+str(e)[:250]
    finally:
        (out/'stop.requested').touch()
        try: process.wait(timeout=15)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid,signal.SIGTERM)
            try: process.wait(timeout=5)
            except subprocess.TimeoutExpired: os.killpg(process.pid,signal.SIGKILL);process.wait()
            record['forced_stop']=True
        record['process_exit']=process.returncode
        save()
record['output_sha256']={p.relative_to(workspace).as_posix():hashlib.sha256(p.read_bytes()).hexdigest()
    for p in workspace.rglob('*') if p.is_file() and not p.is_symlink()}
try:
    if args.case=='local-apps':
        import csv
        def totals(path):
            with sqlite3.connect(path.as_uri()+'?mode=ro',uri=True) as db:
                return list(db.execute('SELECT COUNT(*),SUM(value) FROM records').fetchone())
        with (workspace/'all.csv').open(encoding='utf-8',newline='') as stream:
            rows=list(csv.DictReader(stream))
        note=(workspace/'vault/结果.md').read_text()
        actual={'db':totals(workspace/'records.sqlite'),'backup':totals(workspace/'backup.sqlite'),
                'csv':[len(rows),sum(int(r['value']) for r in rows)],'note_has_link':'[[' in note,
                'graph_exists':(workspace/'links.json').is_file(),'report_exists':(workspace/'report.md').is_file()}
        total=[236,sum(range(235))+10]
        expected={'db':total,'backup':total,'csv':total,'note_has_link':True,'graph_exists':True,'report_exists':True}
    elif args.case=='media':
        actual=json.loads((workspace/'result.json').read_text());expected={'blue_squares':3,'red_triangles':1}
    elif args.case=='local':
        actual=json.loads((workspace/'totals.json').read_text());expected={'杭州':12,'上海':6}
    elif args.case=='browser':
        matches=[p for p in workspace.rglob('result.json') if p.is_file() and not p.is_symlink()]
        artifact=workspace/'result.json' if (workspace/'result.json').is_file() else workspace/'downloads/result.json'
        actual=json.loads(artifact.read_text());expected={'total':73}
        record['downloaded_artifact']=artifact.relative_to(workspace).as_posix()
        record['duplicate_downloads_equal']=all(p.read_bytes()==artifact.read_bytes() for p in matches)
    elif args.case=='document':
        text=(workspace/'report.md').read_text()
        actual={'report_created':True, 'contains_known_totals':all(v in text for v in ('SKU-A','SKU-B','7','11','18'))}
        expected={'report_created':True, 'contains_known_totals':True}
    elif args.case=='novel':
        project=workspace/'managed-novel'
        ledger=json.loads((project/'.novel-system/ledger/chapters.json').read_text())
        chapters=list((project/'正文').glob('*.md'))
        actual={'chapter_count':len(chapters),'content':chapters[0].read_text() if len(chapters)==1 else None,
                'ledger_count':len(ledger),'quality':ledger[0].get('content_quality') if ledger else None,
                'report_created':(workspace/'report.md').is_file()}
        expected={'chapter_count':1,'content':'星光熄灭。\n','ledger_count':1,'quality':'unassessed','report_created':True}
    elif args.case=='quality':
        actual={'project_code_executed':(workspace/'must-not-execute.txt').exists(),
                'report_created':(workspace/'report.md').is_file()}
        expected={'project_code_executed':False,'report_created':True}
    else:
        rows=[json.loads(p.read_text()) for p in (out/'remote-objects').glob('*.json')]
        actual=[{k:v[k] for k in ['label','value','version']} for v in rows]
        expected=[{'label':'本体连接验收','value':41,'version':1}]
        if args.case=='mcp-tasks':
            jobs=json.loads((out/'remote-tasks.json').read_text())
            actual={'records':actual,'tasks':sorted(j['status'] for j in jobs.values())}
            expected={'records':expected,'tasks':['cancelled','completed']}
    record['independent_check']={'actual':actual,'expected':expected,'passed':actual==expected,
        'original_inputs_unchanged':all(hashlib.sha256((workspace/n).read_bytes()).hexdigest()==d
            for n,d in record['input_sha256'].items() if n not in ('aggregate.py','records.sqlite'))}
except Exception as e: record['independent_check']={'passed':False,'error':type(e).__name__}
calls=[json.loads(p.read_text()) for p in (out/'model-calls').glob('*/record.json')]
record['model_metrics']={'calls':len(calls),'failed_calls':sum(c.get('ok') is False for c in calls),
    'unfinished_calls':sum(not c.get('complete') for c in calls),
    'unknown_usage_calls':sum(not c.get('usage') for c in calls),
    'tokens':{k:sum(c.get('usage',{}).get(k,0) or 0 for c in calls if c.get('usage'))
        for k in ('prompt_tokens','completion_tokens','total_tokens','prompt_cache_hit_tokens','prompt_cache_miss_tokens')}}
record['task_success']=(record.get('request_status')=='COMPLETED'
    and record['independent_check']['passed']
    and record['independent_check'].get('original_inputs_unchanged') is True)
if args.case=='mcp':
    fixture.shutdown();fixture.server_close()
save()
print(json.dumps({k:record.get(k) for k in ['case','request_status','task_success','wall_seconds','independent_check','model_metrics','driver_error']},ensure_ascii=False),flush=True)

raise SystemExit(0 if record["task_success"] else 1)
