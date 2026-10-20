from __future__ import annotations

import hashlib
import importlib.machinery
import importlib.util
import marshal
from pathlib import Path
import sys
import tempfile
import types
import unittest


ROOT = Path(__file__).resolve().parents[1]
INTERNAL = ROOT / "app" / "backend" / "tiangong-backend" / "_internal"


def _strings(code: types.CodeType) -> set[str]:
    values: set[str] = set()
    for constant in code.co_consts:
        if isinstance(constant, str):
            values.add(constant)
        elif isinstance(constant, types.CodeType):
            values.update(_strings(constant))
    return values


def _load_frozen_delivery():
    path = INTERNAL / "frozen_modules" / "omni_body_skill" / "tools" / "delivery_kernel.pyc"
    loader = importlib.machinery.SourcelessFileLoader(
        "omni_body_skill.tools._frozen_delivery_mirror_test", str(path),
    )
    spec = importlib.util.spec_from_loader(loader.name, loader)
    assert spec is not None
    module = importlib.util.module_from_spec(spec)
    loader.exec_module(module)
    assert Path(module.__file__).resolve() == path.resolve()
    return module


class FrozenOmniBodyMirrorTests(unittest.TestCase):
    def test_frozen_and_legacy_omni_modules_are_identical_and_contain_hardening(self) -> None:
        expectations = {
            "omni_body_skill/tool_contracts.pyc": {
                "bounded_positive_integer",
                " must be a JSON integer from 1 to 3600 seconds",
                "pptx.read",
            },
            "omni_body_skill/tools/skill_router.pyc": {
                "Skill authority origin must be exact loopback port 7184",
                "/api/v1/internal/skills/",
                "skill.step.check requires an execution-bound Skill activation",
                "tiangong-total-gateway",
            },
            "omni_body_skill/tools/delivery_kernel.pyc": {
                "managed_long_document_delivery",
                "miniapp_page_path_unsafe",
                "content_observation_only",
                "unassessed",
                "adversarial_judge",
                "quality.run_tests",
                "structure_only_not_rendered",
            },
            "omni_body_skill/tools/ppt_design.pyc": {
                "tiangong.v3.ppt_design.v1",
                "TGVisual:card:",
                "builtin:fallback",
            },
            "omni_body_skill/tools/omni_body_tool.pyc": {
                "CapabilityUnavailable",
                "pptx.read",
                "required action missing from delivery registry: ",
            },
        }
        for relative, required in expectations.items():
            frozen = INTERNAL / "frozen_modules" / relative
            legacy = INTERNAL / "legacy_pyz_modules" / relative
            self.assertEqual(hashlib.sha256(frozen.read_bytes()).digest(), hashlib.sha256(legacy.read_bytes()).digest())
            payload = frozen.read_bytes()
            self.assertGreaterEqual(len(payload), 16)
            code = marshal.loads(payload[16:])
            self.assertIsInstance(code, types.CodeType)
            strings = _strings(code)
            self.assertTrue(required.issubset(strings), (relative, sorted(required - strings)))
            if relative == "omni_body_skill/tools/delivery_kernel.pyc":
                # #114 retired heuristic PPT verdicts and A0 test execution.
                # Reject the old artifact even if it retains other hardening.
                retired = {"miniapp_test_command_required", "no_meaningful_visuals", "default_placeholder_layout"}
                self.assertFalse(retired & strings, sorted(retired & strings))

    def test_frozen_code_observation_never_executes_supplied_or_discovered_tests(self) -> None:
        from omni_body_skill.tools.omni_body_tool import BodyRuntime, BodyRuntimeConfig

        delivery = _load_frozen_delivery()
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary).resolve()
            runtime = BodyRuntime(BodyRuntimeConfig(workspace=str(workspace), run_id="frozen-code-observation"))
            marker = workspace / "must-not-execute.txt"
            command = f"from pathlib import Path; Path({str(marker)!r}).write_text('unauthorized')"
            (workspace / "one.py").write_text("x = 1\n", encoding="utf-8")
            (workspace / "test_danger.py").write_text(command, encoding="utf-8")
            for args in ({"test_command": [sys.executable, "-c", command]}, {}):
                with self.subTest(supplied_command=bool(args)):
                    result = delivery.handle_delivery_action(runtime, "frozen-observe", "qc.code.delivery_check", ".", args)
                    self.assertTrue(result["success"])
                    self.assertFalse(marker.exists())
                    report = result["result"]
                    self.assertFalse(report["test_execution"]["executed"])
                    self.assertEqual(report["test_execution"]["required_action"], "quality.run_tests")
                    self.assertEqual(report["content_quality"], "unassessed")
                    self.assertNotIn("score", report)
                    self.assertNotIn("acceptance", report)

    def test_frozen_ppt_observes_default_layout_and_only_explicit_slide_constraint(self) -> None:
        from omni_body_skill.tools.omni_body_tool import BodyRuntime, BodyRuntimeConfig
        from pptx import Presentation

        delivery = _load_frozen_delivery()
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary).resolve()
            runtime = BodyRuntime(BodyRuntimeConfig(workspace=str(workspace), run_id="frozen-ppt-observation"))
            presentation = Presentation()
            for index in range(3):
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = f"Slide {index + 1}"
                slide.placeholders[1].text = "Observed text"
            presentation.save(workspace / "deck.pptx")
            result = delivery.handle_delivery_action(runtime, "frozen-ppt", "qc.ppt.delivery_check", "deck.pptx", {})
            self.assertTrue(result["success"])
            report = result["result"]
            self.assertEqual(report["slides"], 3)
            self.assertEqual(report["visual_coverage"], 0)
            self.assertEqual(report["constraint_checks"], [])
            self.assertEqual(report["issues"], [])
            self.assertEqual(report["content_quality"], "unassessed")
            self.assertEqual(report["layout_observation"], "structure_only_not_rendered")
            self.assertNotIn("score", report)
            self.assertNotIn("acceptance", report)
            constrained = delivery.handle_delivery_action(runtime, "frozen-ppt-explicit", "qc.ppt.delivery_check", "deck.pptx", {"min_slides": 4})
            self.assertEqual(constrained["result"]["constraint_checks"], [
                {"field": "slide_count", "minimum": 4, "actual": 3, "satisfied": False},
            ])

    def test_windows_release_rebuilds_all_frozen_overlays_with_cpython_312(self) -> None:
        rebuild = (ROOT / "scripts" / "rebuild_frozen_release_overlays.py").read_text(encoding="utf-8")
        release = (ROOT / "scripts" / "release-common.mjs").read_text(encoding="utf-8")
        for marker in (
            "sys.version_info[:2] != (3, 12)",
            "omni_body_skill/tool_contracts.pyc",
            "v3/execution_kernel/orchestrator.py",
            "v3/execution_kernel/tool_scheduler.py",
            "tiangong_life/permissions.py",
            "PycInvalidationMode.UNCHECKED_HASH",
        ):
            self.assertIn(marker, rebuild)
        self.assertIn('rebuild_frozen_release_overlays.py', release)


if __name__ == "__main__":
    unittest.main()
