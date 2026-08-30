"""P19-R2 M2 preflight migration parity tests.

`_office_content_gaps` moved (mechanically, zero behaviour change) from
``simple_chain/kernel.py`` to ``simple_chain/content_preflight.py``.
These tests lock:
* the kernel re-export resolves to the very same function object;
* behaviour on the original office scenarios is unchanged (shell docx,
  empty xlsx, placeholder rows, missing columns, row-count requests,
  no-attachment no-op);
* the preflight stays a ``list[str]`` local heuristic and never
  fabricates a ``VerificationRecord`` (authority stays with the oracle).
"""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_xlsx(header: list[str], rows: list[list[object]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if header:
        sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class OfficePreflightParityTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _attachment(self, data: bytes, name: str) -> dict[str, str]:
        path = self.root / name
        path.write_bytes(data)
        return {"path": str(path)}

    def test_kernel_reexport_is_same_object(self) -> None:
        import sys

        sys.path.insert(0, "src")
        sys.path.insert(1, "app/backend/tiangong-backend")
        try:
            from v3.simple_chain.content_preflight import _office_content_gaps as direct
            from v3.simple_chain.kernel import _office_content_gaps as via_kernel
        finally:
            sys.path.remove("src")
            sys.path.remove("app/backend/tiangong-backend")
        self.assertIs(via_kernel, direct)

    def test_behaviour_unchanged_on_original_scenarios(self) -> None:
        from v3.simple_chain.content_preflight import _office_content_gaps

        request = "生成数据表，包含姓名、分数两列，至少5行数据，各3条要点"

        # no attachments -> no gaps
        self.assertEqual(_office_content_gaps(request, []), [])
        self.assertEqual(_office_content_gaps(request, None), [])

        # shell docx (title only) -> gap
        shell = self._attachment(make_docx(["标题"]), "shell.docx")
        gaps = _office_content_gaps(request, [shell])
        self.assertEqual(len(gaps), 1)
        self.assertIn("空壳", gaps[0])

        # filled docx -> no gap
        filled = self._attachment(
            make_docx(
                [
                    "要点一：这一条要点包含足够长的正文内容以满足字数下限要求。",
                    "要点二：这一条同样写入实际内容而不是只有标题占位。",
                    "要点三：第三条要点也有真实内容，凑足全部要求。",
                ]
            ),
            "filled.docx",
        )
        self.assertEqual(_office_content_gaps(request, [filled]), [])

        # empty xlsx (single empty cell) -> 空表 gap
        empty = self._attachment(make_xlsx([], []), "empty.xlsx")
        gaps = _office_content_gaps(request, [empty])
        self.assertEqual(len(gaps), 1)
        self.assertIn("空表", gaps[0])

        # placeholder rows (all identical) -> gap
        placeholder = self._attachment(
            make_xlsx(["姓名", "分数"], [["示例", 0], ["示例", 0]]),
            "placeholder.xlsx",
        )
        gaps = _office_content_gaps(request, [placeholder])
        self.assertTrue(any("全部相同" in gap for gap in gaps), gaps)

        # header + 2 rows while "5行" requested -> row gap (legacy
        # max_row semantics preserved — header counted, by design)
        short = self._attachment(
            make_xlsx(["姓名", "分数"], [["甲", 1], ["乙", 2]]), "short.xlsx"
        )
        gaps = _office_content_gaps(request, [short])
        self.assertTrue(any("行数据" in gap for gap in gaps), gaps)

        # missing column -> gap
        missing = self._attachment(
            make_xlsx(["姓名", "等级"], [["甲", "优"]]), "missing.xlsx"
        )
        gaps = _office_content_gaps(request, [missing])
        self.assertTrue(any("缺少列" in gap and "分数" in gap for gap in gaps), gaps)

    def test_preflight_stays_local_heuristic(self) -> None:
        """The preflight must not fabricate verification-plane objects."""
        import inspect

        from v3.simple_chain import content_preflight

        source = inspect.getsource(content_preflight)
        # 分权边界看导入关系（docstring 提及权威模块名不构成耦合）
        for forbidden in (
            "from contracts.verification import",
            "from total_gateway",
            "import total_gateway",
        ):
            self.assertNotIn(forbidden, source)


if __name__ == "__main__":
    unittest.main()
