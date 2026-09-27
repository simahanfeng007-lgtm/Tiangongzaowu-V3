"""Content observations cannot manufacture quality, completion or execution authority."""
import json
import platform
import sys

import pytest
from omni_body_skill.tools.omni_body_tool import BodyRuntime, BodyRuntimeConfig


@pytest.fixture
def runtime(tmp_path):
    return BodyRuntime(BodyRuntimeConfig(workspace=str(tmp_path), run_id="quality-observation"))


def test_read_only_code_observation_never_executes_supplied_or_discovered_tests(runtime, tmp_path):
    (tmp_path / "one.py").write_text("x = 1\n")
    marker = tmp_path / "must-not-execute.txt"
    code = f"from pathlib import Path; Path({str(marker)!r}).write_text('unauthorized')"
    (tmp_path / "test_danger.py").write_text(code)
    result = runtime.run("qc.code.delivery_check", ".", {"test_command": [sys.executable, "-c", code]})
    assert result["success"] and not marker.exists()
    report = result["result"]
    assert report["test_execution"]["executed"] is False
    assert report["test_execution"]["required_action"] == "quality.run_tests"
    assert "score" not in report and "acceptance" not in report
    assert report["syntax_errors"] == []
    (tmp_path / "broken.py").write_text("def broken(\n")
    broken = runtime.run("qc.code.delivery_check", ".", {})["result"]
    assert len(broken["syntax_errors"]) == 1
    assert broken["content_quality"] == "unassessed"


def test_small_solid_image_is_observed_without_inventing_a_quality_verdict(runtime, tmp_path):
    from PIL import Image
    Image.new("RGB", (3, 5), "white").save(tmp_path / "solid.png")
    for action in ["qc.image.delivery_check", "qc.poster.commercial_check"]:
        result = runtime.run(action, "solid.png", {})
        assert result["success"]
        report = result["result"]
        assert report["image_info"]["luminance_stddev"] == 0
        assert report["visual_content"] == "not_observed"
        assert report["content_quality"] == "unassessed" and "acceptance" not in report
    explicit = runtime.run("qc.image.delivery_check", "solid.png", {"min_width": 4})
    assert explicit["result"]["constraint_checks"] == [{"field": "width", "minimum": 4, "actual": 3, "satisfied": False}]
    (tmp_path / "bad.png").write_bytes(b"broken image")
    assert runtime.run("qc.image.delivery_check", "bad.png", {})["success"] is False


def test_short_deck_without_cta_is_not_blocked_by_default_style_policy(runtime, tmp_path):
    from pptx import Presentation
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(tmp_path / "deck.pptx")
    result = runtime.run("qc.ppt.delivery_check", "deck.pptx", {})
    assert result["success"] and result["result"]["slides"] == 1
    assert result["result"]["constraint_checks"] == []
    assert result["result"]["content_quality"] == "unassessed"
    assert "score" not in result["result"] and "acceptance" not in result["result"]
    result = runtime.run("qc.ppt.delivery_check", "deck.pptx", {"min_slides": 2})
    assert result["result"]["constraint_checks"][0]["satisfied"] is False


def test_sheet_observation_marks_range_and_does_not_claim_formula_or_duplicate_errors(runtime, tmp_path):
    (tmp_path / "values.csv").write_text("value,note\n" + "3,\n" * 1100, encoding="utf-8-sig")
    result = runtime.run("qc.sheet.delivery_check", "values.csv", {})
    assert result["success"]
    report = result["result"]
    assert report["truncated"] and report["rows"] == 1000
    assert report["duplicate_rows_in_observed_range"] == 998
    assert report["blank_cells_in_observed_range"] == 999
    assert report["formula_evaluation"] == "not_performed"
    assert "score" not in report and "acceptance" not in report


def test_native_export_never_substitutes_a_generated_script_for_execution(runtime, tmp_path):
    from omni_body_skill.tools.pro_apps_v34 import _office_native_action
    result = _office_native_action(runtime, "microsoft.word.native.export_pdf", "absent.docx", {})
    assert not result["success"] and result["execution_state"] == "not_executed"
    assert not list(tmp_path.glob("bridges/*"))
    result = runtime.run("microsoft.excel.native.chart.create", "book.xlsx", {})
    assert not result["success"]


def test_local_download_has_no_fabricated_http_status(runtime, tmp_path):
    result = runtime.run("web.download", "data:application/json,%7B%22n%22%3A41%7D", {"output": "data.json"})
    assert result["success"] and result["status"] is None
    assert json.loads((tmp_path / "data.json").read_text()) == {"n": 41}


@pytest.mark.parametrize("suffix", ["docx", "pptx", "xlsx", "pdf", "txt", "png"])
def test_document_preview_never_turns_parse_failure_into_empty_success(runtime, tmp_path, suffix):
    (tmp_path / ("corrupt." + suffix)).write_bytes(b"\xffnot a valid document")
    for action in ("preview.generate", "qc.docx.delivery_check"):
        result = runtime.run(action, "corrupt." + suffix, {})
        assert result["success"] is False


@pytest.mark.parametrize("maximum", [10, 1000, 12000])
def test_preview_reports_source_prefix_and_returned_excerpt_separately(runtime, tmp_path, maximum):
    (tmp_path / "long.txt").write_text("正" * 13000)
    result = runtime.run("qc.docx.delivery_check", "long.txt", {"max_chars": maximum})
    report = result["result"]
    assert result["success"] and report["text_truncated"] and report["preview_truncated"]
    assert report["total_extracted_text_chars"] is None
    assert report["observed_char_range"] == [0, maximum]
    assert report["preview_char_range"] == [0, min(maximum, 1500)]
    assert report["rendered"] is False


def test_real_docx_text_observation_explicitly_excludes_layout(runtime, tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("实际正文")
    document.sections[0].header.paragraphs[0].text = "页眉不在此观察范围"
    document.save(tmp_path / "content.docx")
    result = runtime.run("qc.docx.delivery_check", "content.docx", {})
    assert result["success"] and result["result"]["text_preview"] == "实际正文"
    assert "headers_footnotes_images_layout_not_observed" in result["result"]["observation_scope"]
    assert not result["result"]["text_truncated"]
